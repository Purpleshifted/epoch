"""Builds docs/progress-deck.pptx (editable, 16:9) from the contents below.

Run:  python3 tools/deck/build_deck.py
Needs python-pptx (pip install python-pptx).
Content mirrors docs/*.md; edit the SLIDES section, not the generated file.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

FONT = "Apple SD Gothic Neo"
BG = RGBColor(0x0A, 0x0A, 0x12)
PANEL = RGBColor(0x16, 0x16, 0x22)
PANEL2 = RGBColor(0x1E, 0x1E, 0x2E)
TEXT = RGBColor(0xE8, 0xE8, 0xEE)
DIM = RGBColor(0x9A, 0x9A, 0xAE)
CYAN = RGBColor(0x55, 0xFF, 0xFF)
YELLOW = RGBColor(0xFF, 0xFF, 0x55)
RED = RGBColor(0xFF, 0x55, 0x55)
GREEN = RGBColor(0x55, 0xFF, 0x55)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333
MX = 0.7  # horizontal margin


def rgb_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def textbox(slide, x, y, w, h, paras, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    """paras: list of str or (str, dict(color,size,bold))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, p in enumerate(paras):
        opts = {}
        if isinstance(p, tuple):
            p, opts = p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        para.space_after = Pt(opts.get("space", 6))
        r = para.add_run()
        r.text = p
        r.font.name = FONT
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", bold)
        r.font.color.rgb = opts.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=TEXT):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            txt, opts = it
            paras.append(("•  " + txt, opts))
        else:
            paras.append(("•  " + it, {}))
    return textbox(slide, x, y, w, h, paras, size=size, color=color)


def panel(slide, x, y, w, h, color=PANEL):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.04
    rgb_fill(s, color)
    return s


def new_slide(title, kicker=None, n=None):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if kicker:
        textbox(s, MX, 0.38, 9, 0.35, [(kicker, {"color": CYAN, "size": 12, "bold": True})])
    textbox(s, MX, 0.68, W - 2 * MX, 0.8, [(title, {"size": 28, "bold": True})], anchor=MSO_ANCHOR.TOP)
    if n is not None:
        textbox(s, W - MX - 1, 7.0, 1, 0.3, [(str(n), {"color": DIM, "size": 10})], align=PP_ALIGN.RIGHT)
    return s


def footnote(slide, text):
    textbox(slide, MX, 6.95, W - 2 * MX - 1.2, 0.4, [(text, {"color": DIM, "size": 10})])


def table(slide, x, y, w, rows, col_w, size=12, row_h=0.42, header=True, color_col=None):
    nrows, ncols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = shape.table
    tot = sum(col_w)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Emu(int(Inches(w) * cw / tot))
    for r in range(nrows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            val = rows[r][c]
            col = TEXT
            if isinstance(val, tuple):
                val, col = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL2 if (header and r == 0) else (PANEL if r % 2 else BG)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = val
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = header and r == 0
            run.font.color.rgb = CYAN if (header and r == 0) else col
    return tbl


def flow(slide, y, labels, h=1.05, gap=0.32, size=13, colors=None):
    n = len(labels)
    total_w = W - 2 * MX
    bw = (total_w - gap * (n - 1)) / n
    for i, lab in enumerate(labels):
        x = MX + i * (bw + gap)
        b = panel(slide, x, y, bw, h, PANEL2)
        tf = b.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.08)
        lines = lab.split("\n")
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = line
            r.font.name = FONT
            r.font.size = Pt(size if j == 0 else size - 2)
            r.font.bold = j == 0
            r.font.color.rgb = (colors[i] if colors else CYAN) if j == 0 else DIM
        if i < n - 1:
            textbox(slide, x + bw, y + h / 2 - 0.2, gap, 0.4, [("›", {"color": DIM, "size": 22, "bold": True})],
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def card(slide, x, y, w, h, head, items, head_color=CYAN, size=14):
    panel(slide, x, y, w, h)
    textbox(slide, x + 0.2, y + 0.12, w - 0.4, 0.4, [(head, {"color": head_color, "size": 15, "bold": True})])
    bullets(slide, x + 0.2, y + 0.58, w - 0.4, h - 0.7, items, size=size)


# ───────────────────────── SLIDES ─────────────────────────
n = 0

# 1 ─ title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG
textbox(s, MX, 2.2, 11.5, 0.5, [("PROGRESS REPORT · 2026-10", {"color": CYAN, "size": 14, "bold": True})])
textbox(s, MX, 2.75, 11.8, 1.6, [("Anthropocene\n테크노포실 지층", {"size": 48, "bold": True})], line_spacing=1.05)
textbox(s, MX, 4.75, 11.5, 1.0, [
    ("방문자가 돌아다니며 흩뿌린 것이, 시간이 지나 지층이 된다.", {"size": 20, "color": TEXT}),
    ("피드백 반영 · 학술 근거 · 라이브러리 구축 · 이미지 파이프라인 계획", {"size": 14, "color": DIM}),
])

# 2 ─ direction
n += 1
s = new_slide("피드백을 받아 방향을 바꿨다", "01 · DIRECTION", n + 1)
card(s, MX, 1.7, 5.7, 3.3, "받은 피드백", [
    "별은 좋은 비유가 아니다",
    "방문자는 돌아다니며 동물·플라스틱·전자부품을 흩뿌린다 — 테크노포실 작업",
    "시간이 지나면 지층·미래의 탑뷰에는 오래가는 것이 보인다",
    "\"내가 무엇을 남길지는 내가 정할 수 없고, 주어진 수명 동안 돌아다녀볼 뿐\"",
], head_color=YELLOW)
card(s, MX + 6.0, 1.7, 5.93, 3.3, "반영한 것", [
    "별·중력·병합 물리를 걷어내고 '흩뿌림 + 운명' 구조로 교체 (옛 코드는 archive 브랜치에 보존)",
    "배경 별, 태양풍, 파도 제거 — 움직임의 신호는 스프라이트뿐",
    "방문자가 무엇을 뿌릴지는 추첨 (고를 수 없음)",
    "오래가는 재료만 지층에 남는다",
])
panel(s, MX, 5.25, W - 2 * MX, 1.45, PANEL2)
textbox(s, MX + 0.25, 5.35, W - 2 * MX - 0.5, 1.3, [
    ("그 뒤에 내린 결정", {"color": CYAN, "bold": True, "size": 14}),
    ("과거(깊은 시간의 지표 화석)는 구현에서 뺀다. 지층에는 방문자가 남긴 것만 쌓인다.", {"size": 14}),
    ("인위적 유기물(뼈·조개·숯·음식)은 포함한다. 원래 있던 자연물은 제외하지 않고, 사람이 돌아다닐수록 줄어드는 '자연 스톡'으로 둔다.", {"size": 14}),
])

# 3 ─ system
n += 1
s = new_slide("시스템 구조: 뿌리고, 운명이 정해지고, 쌓인다", "02 · SYSTEM", n + 1)
flow(s, 1.8, [
    "방문자 이동\n걸음 · 정지",
    "무엇을 뿌릴지 추첨\n이웃 땅의 영향",
    "물건의 운명\n분해 · 매몰 · 화석화",
    "지층 / 탑뷰\n오래가는 것만",
], h=1.2, size=15)
card(s, MX, 3.4, 3.8, 3.2, "추첨", [
    "레이어(포장재·전자·유기물)의 사전 확률 × 품목 빈도",
    "주변 반경 안에 살아남은 물건이 확률을 이동시킴 (Pólya urn)",
    "→ 앞사람이 만든 땅을 물려받음",
], size=13)
card(s, MX + 4.05, 3.4, 3.8, 3.2, "운명", [
    "재료마다 수명이 다름 (로그균등)",
    "위에 쌓이면 분해가 50배 느려짐",
    "묻힌 채 시간이 지나면 화석화",
    "소프트 유기물은 1초 안에 사라짐",
], size=13)
card(s, MX + 8.1, 3.4, 3.83, 3.2, "자연 스톡", [
    "칸마다 식물·동물 0~2개",
    "걸음과 쌓인 물건이 '마모 하중'이 됨",
    "하중이 임계값을 넘으면 밀려나 죽은 물질이 됨. 재생 없음",
], size=13, head_color=GREEN)

# 4 ─ catalogue
n += 1
s = new_slide("카탈로그 v2: 흩뿌려지는 품목 148종, 스프라이트 290개", "03 · CATALOGUE", n + 1)
table(s, MX, 1.7, W - 2 * MX, [
    ["레이어", "내용", "품목", "스프라이트", "장면 내 비율"],
    ["A  포장재·플라스틱", "JRC 2016 유럽 해변 쓰레기 표 기반 (측정 빈도)", "106", "178", "70%"],
    ["B  전자제품", "기기·기판 등 (직접 구성)", "32", "53", "10%"],
    ["C  인위적 유기물", "음식물, 배설물, 뼈, 껍데기, 숯 등", "10", "29", "20%"],
    ["N  자연 스톡", "식물·동물. 흩뿌려지지 않고 줄어듦", "18", "30", "—"],
    [("합계", YELLOW), "", ("148 + 18", YELLOW), ("290", YELLOW), ""],
], [2.4, 5.2, 1.1, 1.4, 1.5], size=14, row_h=0.55)
bullets(s, MX, 5.2, W - 2 * MX, 1.6, [
    "품목 하나에 스프라이트 여러 개 — 같은 물건도 다양하게 보이도록. 재료가 아니라 '품목' 단위로 그림",
    "레이어 비율 70/10/20은 내가 제안한 값이며 아직 확정이 아님",
    "수치는 모두 추정(EST.), 재료별 지속성은 순위만 문헌 기반",
], size=14)
footnote(s, "소스: docs/catalogue-v2.md, docs/data/catalogue-v2.json (생성기: tools/catalogue/gen_catalogue_v2.py)")

# 5 ─ evidence
n += 1
s = new_slide("학술 근거와 확인 수준", "04 · EVIDENCE", n + 1)
table(s, MX, 1.7, W - 2 * MX, [
    ["설계 요소", "근거", "확인 수준"],
    ["품목 빈도 (A)", "JRC 2016 해변 쓰레기 Table A1 (유럽)", ("PDF에서 표를 직접 추출해 사용", GREEN)],
    ["동물·식물 감소 비율", "Bar-On et al. 2018, PNAS: 식물 ≈ 1/2, 야생 포유류 ≈ 1/6", ("전문 확인 (PMC)", GREEN)],
    ["식물 마모 곡선", "Cole 1995, J. Applied Ecology: 약 50회 통과 ≈ 식생 50% 감소", ("서지만 확인, 수치는 검색 요약 [S]", YELLOW)],
    ["재료 지속성·화석화 순위", "테크노포실 문헌 (Zalasiewicz 외)", ("순위만 문헌 기반, 수치는 추정", YELLOW)],
    ["이웃 영향 추첨", "Pólya urn", ("근거 문서 작성 (inheritance-rationale)", GREEN)],
], [2.4, 6.0, 3.5], size=13, row_h=0.62)
bullets(s, MX, 5.6, W - 2 * MX, 1.2, [
    "자연 스톡 보정: 기준 하중(칸당 약 50회 통과)에서 식물 생존 59%, 동물 17% — 테스트로 검증",
    "Bar-On 저자들도 '인류 이전 값은 거친 첫 근사'라고 밝힘",
], size=13)
footnote(s, "[S] = 검색 요약만 확인. 검색 요약은 원문과 다른 경우가 있어 라이브 확인이 필요함")

# 6 ─ library
n += 1
s = new_slide("코드와 문서로 구축한 것", "05 · BUILT", n + 1)
card(s, MX, 1.7, 5.9, 2.45, "lib/stratum (결정론적 라이브러리)", [
    "재료 26종, 품목 148 + 자연 18",
    "추첨 · 운명 · 자연 스톡 · 로그 시계",
    "GroundStore: 칸별 타임라인 (증분 계산)",
], size=13)
card(s, MX + 6.1, 1.7, 5.83, 2.45, "검증", [
    "테스트 63개 통과",
    "카탈로그 생성기를 다시 돌려도 원본 데이터가 변하지 않음",
    "tsc: 새 오류 없음 (기존 11건은 SideView·TimespaceView)",
], size=13, head_color=GREEN)
card(s, MX, 4.35, 5.9, 2.35, "문서", [
    "scatter-catalogue / catalogue-v2",
    "inheritance-rationale (이웃 영향의 근거)",
    "performance (비용 모델 + 서버 설계)",
    "sprite-sources-and-pipeline",
], size=13, head_color=YELLOW)
card(s, MX + 6.1, 4.35, 5.83, 2.35, "정리한 것", [
    "별 · 중력 · 병합 · 배경 별 제거",
    "태양풍 · 파도 · NOAA API 제거 (바닥은 평평)",
    "옛 별 코드는 archive/spacetime-stars 에 보존",
], size=13, head_color=RED)

# 7 ─ perf
n += 1
s = new_slide("성능 대응과 전시 서버 설계", "06 · PERFORMANCE", n + 1)
card(s, MX, 1.7, 5.9, 3.0, "지금까지 한 것", [
    "전체 재계산 → 내 주변(반경 24)만 0.5초마다 증분 계산",
    "localStorage 쓰기를 3초마다·변경 시에만",
    "Node 측정: 전체 계산 약 7ms → 반경 계산 약 5ms",
], size=13)
card(s, MX + 6.1, 1.7, 5.83, 3.0, "전시 서버 (미구현 설계)", [
    "서버가 권위를 갖고 칸별 타임라인 보관",
    "폰에는 주변 칸의 변경분만 전송",
    "탑뷰·사이드뷰에는 집계값만",
    "사라진 물건도 시간 기록은 유지 (다른 물건을 묻고 마모시킴)",
    "모든 결과가 해시에서 나오므로 클라이언트가 재계산 가능",
], size=13, head_color=YELLOW)
panel(s, MX, 4.95, W - 2 * MX, 1.7, PANEL2)
textbox(s, MX + 0.25, 5.05, W - 2 * MX - 0.5, 1.5, [
    ("한계", {"color": RED, "bold": True, "size": 14}),
    ("측정은 데스크톱 Node에서만 했고 폰에서는 확인하지 못했다.", {"size": 14}),
    ("버퍼링이 남는다면 원인은 Bloom·Vignette·플레이어 오브 등 다른 곳일 수 있다 (미확인).", {"size": 14}),
])

# 8 ─ static libs
n += 1
s = new_slide("이미지 소스 조사 ① 정적 라이브러리 — 권장 조합", "07 · IMAGE SOURCES", n + 1)
table(s, MX, 1.65, W - 2 * MX, [
    ["소스", "라이선스", "잠재량", "비고"],
    ["Quick, Draw!", "CC BY 4.0", "300+", "28×28 원본. 같은 물건의 변형이 수천 개 — 다양성에 최적"],
    ["game-icons.net", "CC BY 3.0 (일부 CC0)", "300–400", "4,133 SVG. 작가별 크레딧"],
    ["DCSS tiles", "CC0 (미승인 목록 제외)", "150–300", "32×32 픽셀. 음식·동물·뼈"],
    ["Fluent + Noto + Twemoji", "MIT / Apache / CC BY", "스타일당 150–250", "같은 물건을 3가지 스타일로"],
    ["Smithsonian Open Access", "CC0 (표시 항목) [S]", "200–500", "옛 전자제품, 조개·뼈 실물 사진"],
    ["TrashNet", "MIT", "100–300", "흰 판 위 폐기물. 배경 제거 쉬움"],
    ["Openclipart + Health Icons", "CC0 [S] / MIT", "150–300", "범위 넓음. 출처 점검 필요"],
    ["3D 렌더 (GSO, ABO, Poly Haven 등)", "CC BY 4.0 / CC0", "수천", "여러 각도 렌더로 변형 확보"],
], [3.0, 2.6, 1.8, 4.6], size=12, row_h=0.5)
bullets(s, MX, 6.25, W - 2 * MX, 0.7, [
    "1~4번만으로도 목표 290개에 근접. 크레딧은 빌드 때 CREDITS.md 자동 생성",
], size=13)
footnote(s, "조사 환경에서 GitHub 외 다수 사이트가 접근 차단 → [S]는 검색 요약만 확인. 라이브 재확인 항목은 문서 §4에 정리")

# 9 ─ pipeline
n += 1
s = new_slide("이미지 소스 조사 ② 검색어 → 최신 이미지 → 픽셀 스프라이트", "08 · PIPELINE PLAN", n + 1)
flow(s, 1.7, [
    "스케줄러\nGitHub Actions",
    "검색\nCommons 우선",
    "필터\n사람·글자·로고",
    "배경 제거\nrembg",
    "16색 스프라이트\nPillow",
    "저장\n+ 출처 JSON",
], h=1.15, size=12, gap=0.22)
card(s, MX, 3.2, 3.8, 3.5, "필터 / 품질", [
    "사람·얼굴, 텍스트, 워터마크, 로고, NSFW 제외",
    "피사체 면적 15–85%",
    "CLIP 유사도로 '그 물건이 맞는지' 확인",
    "원본·스프라이트 해시로 중복 제거",
], size=12)
card(s, MX + 4.05, 3.2, 3.8, 3.5, "시각 일관성", [
    "32×32 고정, 정사각 크롭",
    "TempleOS 16색, 디더 없음",
    "1px 검정 외곽선",
    "양자화 전 자동 대비·채도 보정",
], size=12, head_color=YELLOW)
card(s, MX + 8.1, 3.2, 3.83, 3.5, "저장 / 버전", [
    "저장소 파일 + atlas.json (선택: R2 미러)",
    "epoch 단위로 쌓음 — 옛 이미지가 퇴적층처럼 남음",
    "원본은 보관하지 않음. 출처 JSON만",
    "삭제 요청 경로 + 라이선스 재검증",
], size=12, head_color=GREEN)
footnote(s, "Pillow 샘플 코드는 아직 실행해 보지 못함. 파이프라인은 계획 단계이며 구현 전")

# 10 ─ constraints
n += 1
s = new_slide("찾아본 것들의 제약", "09 · CONSTRAINTS", n + 1)
table(s, MX, 1.6, W - 2 * MX, [
    ["대상", "상태 / 제약", "판단"],
    ["Google Custom Search", "신규 가입 차단, 기존 고객도 2027-01-01 종료", ("불가", RED)],
    ["Bing Image Search", "2025-08 종료", ("불가", RED)],
    ["Brave Search API", "유료. 저장·ML 이용은 'Storage Rights' 별도 계약, 날짜순 정렬 없음", ("조건부", YELLOW)],
    ["SerpApi 등 스크래핑 중개", "Google이 DMCA 소송 중. 이미지 저작권은 어차피 미해결", ("회피", RED)],
    ["Wikimedia Commons", "키 불필요, 최신순 가능, 파일별 라이선스 표기", ("1순위", GREEN)],
    ["Openverse", "정렬 파라미터가 unstable, '색인일' 기준 (업로드일 아님)", ("보조", YELLOW)],
    ["Flickr", "'합리적 기간만 캐시' 취지의 약관 → 영구 보관과 충돌 소지", ("약관 확인", YELLOW)],
    ["iNaturalist / GBIF", "생물 사진만 (병·휴대폰은 불가)", ("보조", YELLOW)],
    ["Unsplash / Pexels", "저장 금지·데이터셋 금지", ("제외", RED)],
    ["SA·NC·ND 이미지", "ShareAlike는 결과물에 같은 라이선스를 강제", ("제외", RED)],
], [2.8, 6.8, 1.6], size=12, row_h=0.48)
footnote(s, "대부분 검색 요약 기반 [S]. 법률 자문 아님")

# 11 ─ legal
n += 1
s = new_slide("법적 쟁점: 픽셀화해도 원본의 개작이다", "10 · LEGAL", n + 1)
card(s, MX, 1.7, 5.9, 3.2, "비교적 분명한 것", [
    "축소·픽셀화한 복사본도 복제/개작 → CC BY는 크레딧·변경 표기 필요",
    "CC0 / 퍼블릭 도메인은 허락 불필요",
    "사람 얼굴(초상권·GDPR), 상표·로고는 별개 문제",
    "공개 GitHub에 올리면 누구나 재배포 가능 → 파일별 라이선스 필요",
], size=13, head_color=GREEN)
card(s, MX + 6.1, 1.7, 5.83, 3.2, "불확실한 것", [
    "32×32 16색이 '복제'인지 — 의존하지 않음",
    "EU TDM 예외는 분석용이지 변환본 전시 권한이 아님",
    "CJEU Pelham II의 패스티시는 '원작과의 대화'가 필요 → 약한 논거 [S]",
    "한국 저작권법 제35조의5 공정이용: 픽셀화 썸네일 판례는 확인 못함",
], size=13, head_color=YELLOW)
panel(s, MX, 5.15, W - 2 * MX, 1.5, PANEL2)
textbox(s, MX + 0.25, 5.25, W - 2 * MX - 0.5, 1.3, [
    ("정정", {"color": RED, "bold": True, "size": 14}),
    ("TACO 데이터셋은 CC BY 4.0이 아니다. 원본 주석을 읽어보니 이미지별로 라이선스가 섞여 있다 (없음 715장, ODbL 466장, 'CC' 319장). 사진을 저장소에 넣지 않는다.", {"size": 14}),
])
footnote(s, "전시 전에 한국·EU 법률 자문을 따로 받아야 함")

# 12 ─ options
n += 1
s = new_slide("두 가지 설계안", "11 · OPTIONS", n + 1)
table(s, MX, 1.7, W - 2 * MX, [
    ["", "A. 자유 이용 이미지만 (권장)", "B. A + Brave 오픈 웹"],
    ["출처", "Commons, Flickr, iNaturalist/GBIF, Openverse (CC0·PDM·CC-BY)", "A + Brave 이미지 검색"],
    ["작품 문장", "\"세계가 가장 최근에 올린 자유 이용 가능한 이미지\"", "\"웹의 최신 이미지\""],
    ["법적 위험", ("낮음", GREEN), ("높음 — 이미지 저작권 미해결 + 저장 권리는 유료 계약", RED)],
    ["개념 충실도", "중상. 드문 검색어는 후보 간격이 며칠~몇 주 → 그 자체가 하나의 진술", "높음. 단 날짜순이 없어 '최신'은 근사값"],
    ["상태", "바로 만들 수 있음", "변호사 검토 전에는 비권장. 전시장 내부 한정 등 완화책 필요"],
], [1.7, 5.3, 5.0], size=12, row_h=0.62)
bullets(s, MX, 6.45, W - 2 * MX, 0.5, [
    "정적 라이브러리만으로 290개를 채울 수 있으므로, 동적 수집은 그 위에 얹는 층으로 볼 수 있음",
], size=14)

# 13 ─ status
n += 1
s = new_slide("진행 현황과 다음 단계", "12 · STATUS", n + 1)
card(s, MX, 1.7, 3.8, 4.95, "완료", [
    "방향 전환, 별·태양풍·파도 제거",
    "카탈로그 v2 (148 + 18 품목, 스프라이트 290 목표)",
    "추첨 · 운명 · 자연 스톡 라이브러리, 테스트 63개",
    "학술 근거 정리 (일부 [S])",
    "성능 개선 1차, 서버 설계안",
    "이미지 소스·API·법적 조사",
], head_color=GREEN, size=13)
card(s, MX + 4.05, 1.7, 3.8, 4.95, "아직 안 됨", [
    "스프라이트 이미지 제작 (0 / 290)",
    "저해상도 렌더 + 16색 팔레트 셰이더",
    "칩 사운드",
    "옆·위·시공간 뷰 재연결 (현재 데이터가 비어 있음)",
    "수명·퇴장 (10초~3분 랜덤) — 보류",
    "이미지 수집 파이프라인 (계획만)",
], head_color=RED, size=13)
card(s, MX + 8.1, 1.7, 3.83, 4.95, "다음 순서 (제안)", [
    "1. 저해상도 렌더 + 팔레트 셰이더",
    "2. 정적 소스로 스프라이트 첫 묶음 제작",
    "3. 크레딧 자동 생성",
    "4. 뷰 재연결",
    "5. 동적 수집 프로토타입 (Commons 한정)",
    "6. 칩 사운드, 수명·퇴장",
], head_color=YELLOW, size=13)

# 14 ─ decisions
n += 1
s = new_slide("미팅에서 정하고 싶은 것", "13 · DECISIONS", n + 1)
bullets(s, MX, 1.7, W - 2 * MX, 5.0, [
    "개념 문장을 \"자유 이용 가능한 최신 이미지\"로 받아들일 수 있는가? 오픈 웹 전체가 꼭 필요한가?",
    "정적 라이브러리를 기본 층, 동적 수집을 그 위의 층으로 가도 되는가?",
    "동적 수집의 범위: 290개 전부인가, 대표 품목 일부인가?",
    "크레딧 표시: 캔버스에는 텍스트를 넣지 않으므로 별도 페이지에만 둘 것인가?",
    "전시 전 한국·EU 법률 자문을 받을 수 있는가?",
    "열린 파라미터: 레이어 비율(A 70 / B 10 / C 20), 자연 스톡 규모, 회차마다 초기화할지 여부",
    "수치를 문헌에 어떻게 보정할지 (지금은 순위만 문헌 기반, 수치는 추정)",
], size=18)

prs.save("docs/progress-deck.pptx")
print("saved docs/progress-deck.pptx,", len(prs.slides), "slides")
