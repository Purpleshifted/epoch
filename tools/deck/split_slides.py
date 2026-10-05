"""Dev helper: split docs/progress-deck.pptx into one-slide files so Quick Look (first slide only) can preview each."""
import copy, os, sys
from pptx import Presentation

SRC = "docs/progress-deck.pptx"
OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
n = len(Presentation(SRC).slides)
for keep in range(n):
    prs = Presentation(SRC)
    ids = prs.slides._sldIdLst
    for i, sld in reversed(list(enumerate(list(ids)))):
        if i != keep:
            prs.part.drop_rel(sld.rId)
            ids.remove(sld)
    prs.save(os.path.join(OUT, f"s{keep + 1:02d}.pptx"))
print(n, "files in", OUT)
